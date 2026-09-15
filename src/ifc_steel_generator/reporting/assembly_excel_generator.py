from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..models import AssemblyParseResult, AssemblyPart, AssemblyRecord
from .styles import NOTE_FILL, TITLE_FILL, WHITE_BOLD, autosize, style_header, style_total


ASSEMBLY_FILL = PatternFill("solid", fgColor="D9E1F2")


def _safe_stem(source: str | Path) -> str:
    return "".join(c if c not in '<>:"/\\|?*' else "_" for c in Path(source).stem)


def structural_output_filename(source: str | Path) -> str:
    return f"{_safe_stem(source)} lista strukturalna.xlsx"


def shipping_output_filename(source: str | Path) -> str:
    return f"{_safe_stem(source)} lista elementów wysyłkowych.xlsx"


def _natural(value: str) -> tuple[object, ...]:
    return tuple(int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value))


def _clean_material(value: str) -> str:
    parts = [part for part in re.split(r"\s*/\s*", value or "") if part and part.upper() != "STEEL"]
    return "/".join(parts)


@dataclass(slots=True)
class PartLine:
    position: str
    quantity: int
    representative: AssemblyPart

    def unit_mass(self, density: float) -> float:
        return self.representative.element.fabrication_mass_kg(density) or 0.0

    @property
    def unit_surface(self) -> float:
        return self.representative.element.fabrication_surface_m2() or 0.0


@dataclass(slots=True)
class AssemblyGroup:
    mark: str
    phase: str
    records: list[AssemblyRecord]

    @property
    def quantity(self) -> int:
        return len(self.records)

    @property
    def representative(self) -> AssemblyRecord:
        return sorted(self.records, key=lambda value: (-len(value.parts), value.ifc_id))[0]

    @property
    def name(self) -> str:
        values = [record.name for record in self.records if record.name]
        return Counter(values).most_common(1)[0][0] if values else ""

    @property
    def length_mm(self) -> float | None:
        return self.representative.length_mm

    def part_lines(self) -> list[PartLine]:
        grouped: dict[str, list[AssemblyPart]] = defaultdict(list)
        for part in self.representative.parts:
            position = part.element.part_position or part.element.tag or f"IFC-{part.element.ifc_id}"
            grouped[position].append(part)
        lines: list[PartLine] = []
        for position, parts in grouped.items():
            signatures = Counter(
                (
                    part.element.designation,
                    _clean_material(part.element.material),
                    round(part.element.fabrication_length_mm or 0.0),
                    part.element.name,
                )
                for part in parts
            )
            chosen = sorted(signatures, key=lambda value: (-signatures[value], _natural(value[0])))[0]
            representative = next(
                part for part in parts
                if (
                    part.element.designation,
                    _clean_material(part.element.material),
                    round(part.element.fabrication_length_mm or 0.0),
                    part.element.name,
                ) == chosen
            )
            lines.append(PartLine(position, len(parts), representative))
        return sorted(lines, key=lambda value: _natural(value.position))

    def unit_mass(self, density: float) -> float:
        return sum(line.unit_mass(density) * line.quantity for line in self.part_lines())

    def unit_surface(self) -> float:
        return sum(line.unit_surface * line.quantity for line in self.part_lines())


def assembly_groups(result: AssemblyParseResult) -> list[AssemblyGroup]:
    grouped: dict[tuple[str, str], list[AssemblyRecord]] = defaultdict(list)
    for record in result.assemblies:
        grouped[(record.mark, record.phase)].append(record)
    return sorted(
        (AssemblyGroup(mark, phase, records) for (mark, phase), records in grouped.items()),
        key=lambda value: (_natural(value.mark), _natural(value.phase)),
    )


def assembly_report_totals(result: AssemblyParseResult, density: float) -> tuple[int, float, float]:
    groups = assembly_groups(result)
    return (
        sum(group.quantity for group in groups),
        sum(group.unit_mass(density) * group.quantity for group in groups),
        sum(group.unit_surface() * group.quantity for group in groups),
    )


class AssemblyExcelGenerator:
    def generate(self, result: AssemblyParseResult, output_dir: str | Path,
                 density: float = 7850.0) -> tuple[Path, Path]:
        target = Path(output_dir)
        target.mkdir(parents=True, exist_ok=True)
        structural = target / structural_output_filename(result.source)
        shipping = target / shipping_output_filename(result.source)
        self._structural_workbook(result, structural, density)
        self._shipping_workbook(result, shipping, density)
        return structural, shipping

    def _structural_workbook(self, result: AssemblyParseResult, target: Path,
                             density: float) -> None:
        wb = self._workbook(("LISTA STRUKTURALNA", "UWAGI"))
        ws = wb["LISTA STRUKTURALNA"]
        ws.append(["LISTA STRUKTURALNA ZESPOŁÓW MONTAŻOWYCH"])
        ws.merge_cells("A1:I1"); ws["A1"].fill = TITLE_FILL; ws["A1"].font = WHITE_BOLD
        ws.append([])
        ws.append([
            "Pozycja", "Szt.", "Profil / nazwa", "Gatunek", "Długość [mm]",
            "Masa jedn. [kg]", "Masa łączna [kg]", "Powierzchnia łączna [m²]", "Uwagi",
        ])
        style_header(ws[3])
        for group in assembly_groups(result):
            unit_mass = group.unit_mass(density)
            unit_surface = group.unit_surface()
            ws.append([
                group.mark, group.quantity, group.name, "", group.length_mm,
                unit_mass, unit_mass * group.quantity,
                unit_surface * group.quantity, "ZESPÓŁ MONTAŻOWY",
            ])
            for cell in ws[ws.max_row]:
                cell.fill = ASSEMBLY_FILL
                cell.font = Font(bold=True)
            for line in group.part_lines():
                element = line.representative.element
                unit_part_mass = line.unit_mass(density)
                ws.append([
                    line.position, line.quantity, element.designation,
                    _clean_material(element.material), element.fabrication_length_mm,
                    unit_part_mass, unit_part_mass * line.quantity,
                    line.unit_surface * line.quantity, element.name,
                ])
        _, total_mass, total_surface = assembly_report_totals(result, density)
        ws.append(["RAZEM", "", "", "", "", "", total_mass, total_surface, ""])
        style_total(ws[ws.max_row])
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:I{max(3, ws.max_row - 1)}"
        self._notes(wb["UWAGI"], result, density, structural=True)
        self._finish(wb, target)

    def _shipping_workbook(self, result: AssemblyParseResult, target: Path,
                           density: float) -> None:
        wb = self._workbook(("LISTA WYSYŁKOWA", "UWAGI"))
        ws = wb["LISTA WYSYŁKOWA"]
        ws.append(["ZBIORCZA LISTA ELEMENTÓW WYSYŁKOWYCH"])
        ws.merge_cells("A1:I1"); ws["A1"].fill = TITLE_FILL; ws["A1"].font = WHITE_BOLD
        ws.append([])
        ws.append([
            "Zespół montażowy", "Szt.", "Nazwa", "Faza", "Długość [mm]",
            "Masa jedn. [kg]", "Masa łączna [kg]",
            "Powierzchnia jedn. [m²]", "Powierzchnia łączna [m²]",
        ])
        style_header(ws[3])
        for group in assembly_groups(result):
            unit_mass = group.unit_mass(density)
            unit_surface = group.unit_surface()
            ws.append([
                group.mark, group.quantity, group.name, group.phase, group.length_mm,
                unit_mass, unit_mass * group.quantity,
                unit_surface, unit_surface * group.quantity,
            ])
        quantity, total_mass, total_surface = assembly_report_totals(result, density)
        ws.append(["RAZEM", quantity, "", "", "", "", total_mass, "", total_surface])
        style_total(ws[ws.max_row])
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:I{max(3, ws.max_row - 1)}"
        self._notes(wb["UWAGI"], result, density, structural=False)
        self._finish(wb, target)

    @staticmethod
    def _workbook(names: tuple[str, ...]) -> Workbook:
        wb = Workbook(); wb.remove(wb.active)
        for name in names:
            ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False
        return wb

    @staticmethod
    def _notes(ws, result: AssemblyParseResult, density: float, structural: bool) -> None:
        ws.append(["INFORMACJE O RAPORCIE"]); ws["A1"].font = Font(bold=True, size=14)
        quantity, total_mass, total_surface = assembly_report_totals(result, density)
        rows = [
            ("Plik IFC", result.source.name),
            ("Typ raportu", "lista strukturalna" if structural else "zbiorcza lista wysyłkowa"),
            ("Gęstość [kg/m³]", density),
            ("Pozycje zespołów", len(assembly_groups(result))),
            ("Fizyczne zespoły", quantity),
            ("Masa łączna [kg]", total_mass),
            ("Powierzchnia łączna [m²]", total_surface),
            ("Reguła struktury", "zespół wg ASSEMBLY_POS; części wg PART_POS"),
            ("Reguła masy", "profile: dokładna masa tabelaryczna × długość warsztatowa; pozostałe: geometria IFC"),
            ("Reguła powierzchni", "zewnętrzna powierzchnia malowana; bez wnętrza profili zamkniętych"),
        ]
        for row in rows:
            ws.append(list(row))
        ws.append([]); ws.append(["OSTRZEŻENIA"]); ws[ws.max_row][0].font = Font(bold=True)
        for warning in result.warnings:
            ws.append([warning]); ws[ws.max_row][0].fill = NOTE_FILL

    @staticmethod
    def _finish(wb: Workbook, target: Path) -> None:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    header = str(ws.cell(3, cell.column).value or "") if ws.title != "UWAGI" else ""
                    if "[kg]" in header:
                        cell.number_format = '0.000 "kg"'
                    elif "[m²]" in header:
                        cell.number_format = '0.000 "m²"'
                    elif "[mm]" in header:
                        cell.number_format = '0.00 "mm"'
            autosize(ws, maximum=42)
        try:
            wb.save(target)
        except PermissionError as exc:
            raise PermissionError(f"Nie można zapisać {target.name}. Zamknij plik w Excelu.") from exc
