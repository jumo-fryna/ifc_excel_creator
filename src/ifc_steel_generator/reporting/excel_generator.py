from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..models import ParseResult, SteelElement
from .styles import NOTE_FILL, TITLE_FILL, WHITE_BOLD, autosize, style_header, style_total

SHEETS = ["PODSUMOWANIE", "PROFILE", "BLACHY", "DANE_PROFILE", "DANE_BLACHY"]


def output_filename(source: str | Path) -> str:
    safe_stem = "".join(c if c not in '<>:"/\\|?*' else "_" for c in Path(source).stem)
    return f"{safe_stem} lista profili analiza.xlsx"


class ExcelGenerator:
    def generate(self, result: ParseResult, output: str | Path, density: float = 7850.0) -> Path:
        target = Path(output)
        target.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)
        for name in SHEETS:
            wb.create_sheet(name)
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.calculation.calcMode = "auto"
        except AttributeError:
            pass
        self._summary(wb["PODSUMOWANIE"], result, density)
        self._profiles(wb["PROFILE"], result.profiles, density)
        self._plates(wb["BLACHY"], result.plates, density)
        self._profile_data(wb["DANE_PROFILE"], result.profiles, density)
        self._plate_data(wb["DANE_BLACHY"], result.plates, density)
        for ws in wb.worksheets:
            self._number_formats(ws)
            autosize(ws)
        try:
            wb.save(target)
        except PermissionError as exc:
            raise PermissionError(f"Nie można zapisać {target.name}. Zamknij plik w Excelu.") from exc
        return target

    @staticmethod
    def _mass(elements: list[SteelElement], density: float) -> float:
        return sum(x.mass_kg(density) or 0 for x in elements)

    @staticmethod
    def _net_mass(elements: list[SteelElement], density: float) -> float:
        return sum(x.net_mass_kg(density) or 0 for x in elements)

    def _summary(self, ws, result: ParseResult, density: float) -> None:
        ws.append(["ZESTAWIENIE STALI Z MODELU IFC"])
        ws.merge_cells("A1:H1")
        ws["A1"].fill = TITLE_FILL; ws["A1"].font = WHITE_BOLD
        ws.append([])
        ws.append([
            "Grupa", "Szt.", "Długość [m]", "NetVolume [m³]",
            "Masa materiałowa [kg]", "Masa materiałowa [t]",
            "Masa netto IFC (kontrolna) [kg]", "Masa netto IFC (kontrolna) [t]",
        ])
        style_header(ws[3])
        pm = self._mass(result.profiles, density); bm = self._mass(result.plates, density)
        pnm = self._net_mass(result.profiles, density); bnm = self._net_mass(result.plates, density)
        plen = sum((p.length_mm or 0) / 1000 for p in result.profiles)
        pv = sum(p.net_volume_m3 or 0 for p in result.profiles); bv = sum(p.net_volume_m3 or 0 for p in result.plates)
        ws.append(["Profile", len(result.profiles), plen, pv, pm, pm / 1000, pnm, pnm / 1000])
        ws.append(["Blachy", len(result.plates), 0, bv, bm, bm / 1000, bnm, bnm / 1000])
        ws.append(["RAZEM", len(result.profiles)+len(result.plates), plen, pv+bv, pm+bm, (pm+bm)/1000, pnm+bnm, (pnm+bnm)/1000])
        style_total(ws[6])
        ws.append(["Gęstość stali [kg/m³]", density])
        ws.append(["Blachy bez grubości w nazwie", result.unknown_plate_thickness])
        ws.append([]); ws.append(["PODSUMOWANIE WG MATERIAŁU"]); ws["A10"].font = Font(bold=True)
        ws.append([
            "Materiał", "Profile [szt.]", "Blachy [szt.]",
            "Masa materiałowa [kg]", "Masa materiałowa [t]",
            "Masa netto IFC (kontrolna) [kg]", "Masa netto IFC (kontrolna) [t]",
        ]); style_header(ws[11])
        materials = sorted({x.material or "brak danych" for x in result.profiles + result.plates})
        for mat in materials:
            ps=[x for x in result.profiles if (x.material or "brak danych")==mat]
            bs=[x for x in result.plates if (x.material or "brak danych")==mat]
            values=ps+bs; mass=self._mass(values,density); net_mass=self._net_mass(values,density)
            ws.append([mat,len(ps),len(bs),mass,mass/1000,net_mass,net_mass/1000])
        row=ws.max_row+2; ws.cell(row,1,"PODSUMOWANIE BLACH WG GRUBOŚCI").font=Font(bold=True)
        row+=1; headers=[
            "Grubość [mm]", "Szt.", "NetArea [m²]", "NetVolume [m³]",
            "Masa materiałowa [kg]", "Masa materiałowa [t]",
            "Masa netto IFC (kontrolna) [kg]", "Masa netto IFC (kontrolna) [t]",
        ]
        for c,v in enumerate(headers,1): ws.cell(row,c,v)
        style_header(ws[row])
        groups: dict[float | None,list[SteelElement]]=defaultdict(list)
        for p in result.plates: groups[p.thickness_mm].append(p)
        for key in sorted(groups, key=lambda x:(x is None,x or 0)):
            values=groups[key]; mass=self._mass(values,density); net_mass=self._net_mass(values,density); row+=1
            ws.append(["brak danych" if key is None else key,len(values),sum(x.net_area_m2 or 0 for x in values),sum(x.net_volume_m3 or 0 for x in values),mass,mass/1000,net_mass,net_mass/1000])
        row=ws.max_row+2
        notes=[
            "Źródło:",
            "Masa materiałowa profili: Weight z IFC; przy braku tabela kg/m, następnie NetWeight i geometria.",
            "Masa materiałowa blach: WeightNet z IFC; przy braku Weight, objętość albo geometria.",
            "Masa netto IFC (kontrolna): WeightNet albo NetVolume × gęstość; nie jest podstawą tonażu, gdy IFC zawiera Weight.",
            "Dla oznaczeń bez pewnej informacji o grubości pola grubości pozostają puste.",
        ]
        for note in notes:
            ws.cell(row,1,note).fill=NOTE_FILL; row+=1
        ws.freeze_panes="A4"

    def _profiles(self, ws, elements: list[SteelElement], density: float) -> None:
        headers=[
            "Profil", "Materiał", "Szt.", "Długość łączna [m]",
            "Masa materiałowa [kg]", "Masa materiałowa [t]", "Średnio [kg/m]",
            "Masa netto IFC (kontrolna) [kg]", "Masa netto IFC (kontrolna) [t]",
        ]
        ws.append(headers); style_header(ws[1]); groups=defaultdict(list)
        for e in elements: groups[(e.designation,e.material)].append(e)
        for (name,mat),values in sorted(groups.items()):
            length=sum((x.length_mm or 0)/1000 for x in values); mass=self._mass(values,density); net_mass=self._net_mass(values,density)
            ws.append([name,mat,len(values),length,mass,mass/1000,mass/length if length else None,net_mass,net_mass/1000])
        mass=self._mass(elements,density); net_mass=self._net_mass(elements,density)
        ws.append(["RAZEM PROFILE","",len(elements),sum((x.length_mm or 0)/1000 for x in elements),mass,mass/1000,None,net_mass,net_mass/1000]); style_total(ws[ws.max_row])
        ws.auto_filter.ref=f"A1:I{max(1,ws.max_row-1)}"; ws.freeze_panes="A2"

    def _plates(self, ws, elements: list[SteelElement], density: float) -> None:
        headers=[
            "Oznaczenie", "Grubość [mm]", "Szer. nom. [mm]", "Materiał", "Szt.",
            "NetArea [m²]", "GrossArea [m²]", "NetVolume [m³]", "GrossVolume [m³]",
            "Masa materiałowa [kg]", "Masa materiałowa [t]",
            "Masa netto IFC (kontrolna) [kg]", "Masa netto IFC (kontrolna) [t]",
        ]
        ws.append(headers); style_header(ws[1]); groups=defaultdict(list)
        for e in elements: groups[(e.designation,e.material)].append(e)
        for (name,mat),values in sorted(groups.items()):
            first=values[0]; mass=self._mass(values,density); net_mass=self._net_mass(values,density)
            ws.append([name,first.thickness_mm,first.nominal_width_mm,mat,len(values),sum(x.net_area_m2 or 0 for x in values),sum(x.gross_area_m2 or 0 for x in values),sum(x.net_volume_m3 or 0 for x in values),sum(x.gross_volume_m3 or 0 for x in values),mass,mass/1000,net_mass,net_mass/1000])
        mass=self._mass(elements,density); net_mass=self._net_mass(elements,density)
        ws.append(["RAZEM BLACHY","","","",len(elements),sum(x.net_area_m2 or 0 for x in elements),sum(x.gross_area_m2 or 0 for x in elements),sum(x.net_volume_m3 or 0 for x in elements),sum(x.gross_volume_m3 or 0 for x in elements),mass,mass/1000,net_mass,net_mass/1000]); style_total(ws[ws.max_row])
        ws.auto_filter.ref=f"A1:M{max(1,ws.max_row-1)}"; ws.freeze_panes="A2"

    def _profile_data(self, ws, elements: list[SteelElement], density: float) -> None:
        headers=[
            "IFC ID", "Typ IFC", "Profil", "Marka/Tag", "Materiał", "Długość [mm]",
            "NetVolume [m³]", "Weight [kg]", "NetWeight [kg]", "Masa tabelaryczna [kg/m]",
            "Źródło masy", "Masa materiałowa [kg]", "Masa netto IFC (kontrolna) [kg]",
            "OuterSurfaceArea [m²]",
        ]
        ws.append(headers); style_header(ws[1])
        for e in elements: ws.append([e.ifc_id,e.ifc_type,e.designation,e.tag,e.material,e.length_mm,e.net_volume_m3,e.gross_weight_kg,e.net_weight_kg,e.unit_weight_kg_m,e.mass_source,e.mass_kg(density),e.net_mass_kg(density),e.outer_surface_area_m2])
        self._table(ws,"DaneProfile")

    def _plate_data(self, ws, elements: list[SteelElement], density: float) -> None:
        headers=[
            "IFC ID", "Oznaczenie blachy", "Marka/Tag", "Materiał", "Grubość [mm]",
            "Szerokość nominalna [mm]", "NetArea [m²]", "GrossArea [m²]", "NetVolume [m³]",
            "GrossVolume [m³]", "Weight [kg]", "NetWeight [kg]", "Źródło masy",
            "Masa materiałowa [kg]", "Masa netto IFC (kontrolna) [kg]",
        ]
        ws.append(headers); style_header(ws[1])
        for e in elements:
            ws.append([e.ifc_id,e.designation,e.tag,e.material,e.thickness_mm,e.nominal_width_mm,e.net_area_m2,e.gross_area_m2,e.net_volume_m3,e.gross_volume_m3,e.gross_weight_kg,e.net_weight_kg,e.mass_source,e.mass_kg(density),e.net_mass_kg(density)])
        self._table(ws,"DaneBlachy")

    @staticmethod
    def _table(ws, name: str) -> None:
        if ws.max_row >= 2:
            table=Table(displayName=name,ref=f"A1:{ws.cell(1,ws.max_column).column_letter}{ws.max_row}")
            table.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
            ws.add_table(table)
        else:
            ws.auto_filter.ref=f"A1:{ws.cell(1,ws.max_column).column_letter}1"
        ws.freeze_panes="A2"

    @staticmethod
    def _number_formats(ws) -> None:
        """Applies formats based on Polish report header semantics."""
        header_rows = []
        for row in ws.iter_rows():
            values = [str(c.value or "") for c in row]
            unit_headers = [
                value for value in values
                if any(unit in value for unit in ("[m³]", "[kg]", "[kg/m]", "[mm]", "[m²]", "[t]", "[m]"))
            ]
            if len(unit_headers) >= 2:
                header_rows.append(row[0].row)
        for index, header_row in enumerate(header_rows):
            end_row = header_rows[index + 1] - 1 if index + 1 < len(header_rows) else ws.max_row
            for cell in ws[header_row]:
                title = str(cell.value or "")
                if "[m³]" in title:
                    fmt = '0.000000 "m³"'
                elif "[m²]" in title:
                    fmt = '0.000 "m²"'
                elif "[kg/m]" in title:
                    fmt = '0.000 "kg/m"'
                elif "[kg]" in title:
                    fmt = '0.00 "kg"'
                elif "[t]" in title:
                    fmt = '0.000 "t"'
                elif "[mm]" in title:
                    fmt = '0.00 "mm"'
                elif "[m]" in title:
                    fmt = '0.000 "m"'
                else:
                    continue
                for row in range(header_row + 1, end_row + 1):
                    ws.cell(row, cell.column).number_format = fmt
