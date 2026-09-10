from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DARK_BLUE = "17365D"
BLUE = "4472C4"
LIGHT_GREEN = "E2F0D9"
LIGHT_YELLOW = "FFF2CC"
WHITE = "FFFFFF"

TITLE_FILL = PatternFill("solid", fgColor=DARK_BLUE)
HEADER_FILL = PatternFill("solid", fgColor=BLUE)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT_GREEN)
NOTE_FILL = PatternFill("solid", fgColor=LIGHT_YELLOW)
WHITE_BOLD = Font(color=WHITE, bold=True)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")


def style_header(row) -> None:
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = CENTER


def style_total(row) -> None:
    for cell in row:
        cell.fill = TOTAL_FILL
        cell.font = BOLD


def autosize(ws, maximum: int = 42) -> None:
    for index, column in enumerate(ws.columns, 1):
        letter = get_column_letter(index)
        width = max((len(str(c.value)) if c.value is not None else 0) for c in column) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), maximum)
