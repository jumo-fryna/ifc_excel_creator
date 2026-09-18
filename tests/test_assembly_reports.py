from pathlib import Path

import ifcopenshell
from openpyxl import load_workbook

from ifc_steel_generator.assembly_parser import IfcAssemblyParser
from ifc_steel_generator.reporting.assembly_excel_generator import (
    AssemblyExcelGenerator,
    shipping_output_filename,
    structural_output_filename,
)


def _shape(model, name: str, x: float, y: float, depth: float):
    origin = model.create_entity("IfcCartesianPoint", Coordinates=(0.0, 0.0, 0.0))
    axis = model.create_entity("IfcAxis2Placement3D", Location=origin)
    direction = model.create_entity("IfcDirection", DirectionRatios=(0.0, 0.0, 1.0))
    p2 = model.create_entity("IfcAxis2Placement2D", Location=model.create_entity("IfcCartesianPoint", Coordinates=(0.0, 0.0)))
    profile = model.create_entity("IfcRectangleProfileDef", ProfileType="AREA", ProfileName=name, Position=p2, XDim=x, YDim=y)
    solid = model.create_entity("IfcExtrudedAreaSolid", SweptArea=profile, Position=axis, ExtrudedDirection=direction, Depth=depth)
    context = model.create_entity("IfcGeometricRepresentationContext", ContextIdentifier="Body", ContextType="Model", CoordinateSpaceDimension=3, Precision=1e-5, WorldCoordinateSystem=axis)
    representation = model.create_entity("IfcShapeRepresentation", ContextOfItems=context, RepresentationIdentifier="Body", RepresentationType="SweptSolid", Items=[solid])
    return model.create_entity("IfcProductDefinitionShape", Representations=[representation])


def _property_set(model, element, values):
    properties = [
        model.create_entity("IfcPropertySingleValue", Name=name, NominalValue=model.create_entity(kind, value))
        for name, kind, value in values
    ]
    pset = model.create_entity("IfcPropertySet", GlobalId=ifcopenshell.guid.new(), Name="Test", HasProperties=properties)
    model.create_entity("IfcRelDefinesByProperties", GlobalId=ifcopenshell.guid.new(), RelatedObjects=[element], RelatingPropertyDefinition=pset)


def _model(path: Path) -> None:
    model = ifcopenshell.file(schema="IFC4")
    mm = model.create_entity("IfcSIUnit", UnitType="LENGTHUNIT", Prefix="MILLI", Name="METRE")
    m2 = model.create_entity("IfcSIUnit", UnitType="AREAUNIT", Name="SQUARE_METRE")
    m3 = model.create_entity("IfcSIUnit", UnitType="VOLUMEUNIT", Name="CUBIC_METRE")
    kg = model.create_entity("IfcSIUnit", UnitType="MASSUNIT", Prefix="KILO", Name="GRAM")
    units = model.create_entity("IfcUnitAssignment", Units=[mm, m2, m3, kg])
    model.create_entity("IfcProject", GlobalId=ifcopenshell.guid.new(), UnitsInContext=units)
    assembly = model.create_entity("IfcElementAssembly", GlobalId=ifcopenshell.guid.new(), Name="Träger", Tag="A-01", PredefinedType="NOTDEFINED")
    plate = model.create_entity("IfcPlate", GlobalId=ifcopenshell.guid.new(), Name="PL10x100", Tag="P1", Representation=_shape(model, "PL10x100", 10.0, 100.0, 1000.0))
    beam = model.create_entity("IfcBeam", GlobalId=ifcopenshell.guid.new(), Name="IPE100", Tag="B1", Representation=_shape(model, "IPE100", 100.0, 10.0, 2000.0))
    model.create_entity("IfcRelAggregates", GlobalId=ifcopenshell.guid.new(), RelatingObject=assembly, RelatedObjects=[plate, beam])
    _property_set(model, assembly, [
        ("ASSEMBLY_POS", "IfcIdentifier", "A-01"),
        ("ShippingMark", "IfcIdentifier", "S-01"),
        ("PHASE", "IfcInteger", 10),
        ("Assembly/Cast unit weight", "IfcMassMeasure", 25.0),
    ])
    path.parent.mkdir(parents=True, exist_ok=True); model.write(str(path))


def _one_piece_model(path: Path) -> None:
    model = ifcopenshell.file(schema="IFC4")
    mm = model.create_entity("IfcSIUnit", UnitType="LENGTHUNIT", Prefix="MILLI", Name="METRE")
    m2 = model.create_entity("IfcSIUnit", UnitType="AREAUNIT", Name="SQUARE_METRE")
    m3 = model.create_entity("IfcSIUnit", UnitType="VOLUMEUNIT", Name="CUBIC_METRE")
    kg = model.create_entity("IfcSIUnit", UnitType="MASSUNIT", Prefix="KILO", Name="GRAM")
    units = model.create_entity("IfcUnitAssignment", Units=[mm, m2, m3, kg])
    model.create_entity("IfcProject", GlobalId=ifcopenshell.guid.new(), UnitsInContext=units)
    for index in range(2):
        beam = model.create_entity(
            "IfcBeam", GlobalId=ifcopenshell.guid.new(), Name="Belka jednoczęściowa",
            Tag=f"A-02-{index}", Representation=_shape(model, "IPE100", 100.0, 10.0, 2000.0),
        )
        _property_set(model, beam, [
            ("ASSEMBLY_POS", "IfcIdentifier", "A-02"),
            ("PART_POS", "IfcIdentifier", "A-02"),
            ("PROFILE", "IfcIdentifier", "IPE100"),
            ("PHASE", "IfcInteger", 10),
        ])
    path.parent.mkdir(parents=True, exist_ok=True); model.write(str(path))


def test_assembly_parser_and_two_workbooks(tmp_path):
    source = tmp_path / "model.ifc"; _model(source)
    result = IfcAssemblyParser().parse(source, phase=("10",))
    assert len(result.assemblies) == 1
    assert result.assemblies[0].mark == "A-01"
    assert result.assemblies[0].shipping_mark == "S-01"
    assert len(result.parts) == 2
    # The report is reconciled from its physical parts; the rounded assembly
    # weight remains diagnostic rather than overriding the bill of materials.
    assert result.assemblies[0].report_mass_kg(7850) == 24.021
    assert result.assemblies[0].surface_area_m2 > 0

    structural, shipping = AssemblyExcelGenerator().generate(result, tmp_path)
    assert structural.name == structural_output_filename(source)
    assert shipping.name == shipping_output_filename(source)
    structural_wb = load_workbook(structural)
    shipping_wb = load_workbook(shipping)
    assert structural_wb.sheetnames == ["LISTA STRUKTURALNA", "UWAGI"]
    assert shipping_wb.sheetnames == ["LISTA WYSYŁKOWA", "UWAGI"]
    rows = list(structural_wb["LISTA STRUKTURALNA"].iter_rows(min_row=4, values_only=True))
    assert rows[0][0:3] == ("A-01", 1, "Träger")
    assert {row[2] for row in rows[1:3]} == {"PL10x100", "IPE100"}
    shipping_rows = list(shipping_wb["LISTA WYSYŁKOWA"].iter_rows(min_row=4, values_only=True))
    assert shipping_rows[0][0:3] == ("A-01", 1, "Träger")


def test_assembly_opens_model_once(tmp_path):
    from unittest.mock import patch
    source = tmp_path / "single-read.ifc"
    _model(source)
    with patch("ifcopenshell.open", wraps=ifcopenshell.open) as opened:
        result = IfcAssemblyParser().parse(source)
    assert opened.call_count == 1
    assert len(result.parts) == 2


def test_one_piece_assemblies_are_detected_without_ifcelementassembly(tmp_path):
    source = tmp_path / "one-piece.ifc"
    _one_piece_model(source)
    result = IfcAssemblyParser().parse(source, phase=("10",))
    assert len(result.assemblies) == 2
    assert {record.mark for record in result.assemblies} == {"A-02"}
    assert all(len(record.parts) == 1 for record in result.assemblies)


def test_unassigned_parts_are_in_both_report_totals_without_invented_assemblies(tmp_path):
    import pytest
    from ifc_steel_generator.reporting.assembly_excel_generator import assembly_report_totals
    source = tmp_path / "unassigned.ifc"
    _one_piece_model(source)
    model = ifcopenshell.open(str(source))
    for relation in model.by_type("IfcRelDefinesByProperties"):
        model.remove(relation)
    model.write(str(source))
    result = IfcAssemblyParser().parse(source)
    assert len(result.assemblies) == 0
    assert len(result.unassigned_elements) == 2
    count, mass, surface = assembly_report_totals(result, 7850)
    assert count == 0
    assert mass == pytest.approx(2 * 2 * 8.0855)
    assert surface > 0
    for path in AssemblyExcelGenerator().generate(result, tmp_path):
        rows = list(load_workbook(path, data_only=True).worksheets[0].values)
        total = next(r for r in rows if r[0] == "RAZEM")
        assert total[6] == pytest.approx(mass)
        assert sum(str(r[0]).startswith(("IFC-", "BRAK ZESPOŁU / IFC-")) for r in rows) == 2


def test_workshop_column_correction_does_not_change_material_parser(tmp_path):
    import pytest
    from unittest.mock import patch
    from types import SimpleNamespace
    from array import array
    from ifc_steel_generator.models import SteelElement, ElementKind
    from ifc_steel_generator.assembly_parser import WorkshopElementParser
    from ifc_steel_generator.units import UnitConverter
    item = SteelElement(9999, "IfcColumn", ElementKind.PROFILE,
                        designation="HEA200", length_mm=690.15,
                        stock_length_mm=690.15, unit_weight_kg_m=42.233)
    material_mass = item.mass_kg(7850)
    vertices = [v for x in (0., .2) for y in (0., .2) for z in (0., .66015) for v in (x,y,z)]
    geometry = SimpleNamespace(verts=vertices, verts_buffer=array('d', vertices).tobytes(), faces=[])
    with patch("ifc_steel_generator.parser.IfcParser._apply_geometry_fallback"):
        WorkshopElementParser()._apply_geometry_fallback(None, item, UnitConverter(), None, geometry)
    assert item.fabrication_length_mm == pytest.approx(660.15)
    assert item.mass_kg(7850) == material_mass
    assert item.fabrication_mass_kg(7850) == pytest.approx(42.233 * .66015)
