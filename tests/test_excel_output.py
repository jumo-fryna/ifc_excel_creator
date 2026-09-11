from pathlib import Path

from openpyxl import load_workbook

from ifc_steel_generator.models import ElementKind, ParseResult, SteelElement
from ifc_steel_generator.reporting.excel_generator import ExcelGenerator, SHEETS, output_filename


def sample():
    return ParseResult(
        Path("model.ifc"),
        profiles=[SteelElement(
            1, "IfcBeam", ElementKind.PROFILE, "HEA300", material="S355",
            length_mm=1000, gross_weight_kg=100.0, net_weight_kg=88.3,
        )],
        plates=[SteelElement(
            2, "IfcPlate", ElementKind.PLATE, "PL16x144", material="S355",
            thickness_mm=16, nominal_width_mm=144, net_volume_m3=0.01,
            gross_weight_kg=80.0,
        )],
    )


def test_workbook(tmp_path):
    target=ExcelGenerator().generate(sample(),tmp_path/"out.xlsx")
    wb=load_workbook(target,data_only=False)
    assert wb.sheetnames==SHEETS
    assert wb["DANE_BLACHY"]["N2"].value==80.0
    assert wb["DANE_BLACHY"]["O2"].value==78.5
    assert wb["PODSUMOWANIE"]["G3"].value=="Masa netto IFC (kontrolna) [kg]"
    assert wb["PODSUMOWANIE"]["E6"].value==180.0
    assert wb["PODSUMOWANIE"]["G6"].value==166.8
    assert wb["PODSUMOWANIE"]["B7"].value==7850
    assert wb["PODSUMOWANIE"]["C12"].number_format=="General"
    assert wb["PROFILE"]["G2"].number_format=='0.000 "kg/m"'


def test_filename():
    assert output_filename("abc-C-003-rev1.ifc")=="abc-C-003-rev1 lista profili analiza.xlsx"
    assert output_filename("abc.ifc")=="abc lista profili analiza.xlsx"
